"""Leitores: transformam o arquivo da aula (.txt, .pdf, .pptx, vídeo/áudio) num `Documento`.

Bibliotecas externas só são importadas quando o tipo de arquivo aparece,
então dá para usar .txt sem instalar nada.
"""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

from .modelo import Bloco, Documento

EXT_TEXTO = {".txt", ".md"}
EXT_PDF = {".pdf"}
EXT_PPT = {".pptx"}
EXT_MIDIA = {".mp4", ".mkv", ".mov", ".webm", ".mp3", ".wav", ".m4a", ".ogg"}

MARCA_LISTA = re.compile(r"^\s*(?:[-•–*▪●○◦‣⁃■□➢✓])\s+(.*)$")
MARCA_NUMERO = re.compile(r"^\s*(?:\d{1,2}|[a-z])[.)]\s+(.*)$", re.I)
NOME_DE_ARQUIVO = re.compile(r".*\.(png|jpe?g|gif|bmp|tiff?|svg|emf|wmf|webp)$", re.I)


def ler(caminho: Path, pasta_trabalho: Path, modelo_audio: str = "small") -> Documento:
    """Escolhe o leitor pela extensão. Imagens extraídas vão para pasta_trabalho/imagens."""
    caminho = Path(caminho)
    ext = caminho.suffix.lower()
    if ext in EXT_TEXTO:
        doc = _ler_texto(caminho)
    elif ext in EXT_PDF:
        doc = _ler_pdf(caminho, pasta_trabalho)
    elif ext in EXT_PPT:
        doc = _ler_pptx(caminho, pasta_trabalho)
    elif ext in EXT_MIDIA:
        doc = _ler_midia(caminho, modelo_audio)
    else:
        raise ValueError(
            f"Formato não suportado: {ext}. Use .pdf, .pptx, .txt, .mp4 ou .mp3."
        )
    doc.origem = caminho.name
    return doc


# ---------- utilidades ----------

def _tirar_titulo(blocos: list[Bloco], padrao: str) -> tuple[str, list[Bloco]]:
    """O título principal (nível mais alto entre os 3 primeiros blocos) vira o título do documento.
    Cobre apostilas com cabeçalho pequeno antes do título ("Ciências - 5º ano")."""
    candidatos = [(b.nivel, i) for i, b in enumerate(blocos[:3]) if b.tipo == "titulo"]
    if not candidatos:
        return padrao, blocos
    _, i = min(candidatos)
    return blocos[i].texto, blocos[:i] + blocos[i + 1:]


def _bloco_de_linhas(linhas: list[str]) -> Bloco | None:
    """Decide se um grupo de linhas é lista, etapas ou parágrafo."""
    linhas = [l for l in (x.strip() for x in linhas) if l]
    if not linhas:
        return None
    if all(MARCA_NUMERO.match(l) for l in linhas) and len(linhas) >= 2:
        return Bloco("etapas", itens=[MARCA_NUMERO.match(l).group(1) for l in linhas])
    if all(MARCA_LISTA.match(l) for l in linhas):
        return Bloco("lista", itens=[MARCA_LISTA.match(l).group(1) for l in linhas])
    return Bloco("paragrafo", texto=" ".join(linhas))


def _parece_titulo(linha: str) -> bool:
    linha = linha.strip()
    return (
        0 < len(linha) <= 70
        and not linha.endswith((".", ":", ";", ",", "!", "?"))
        and not MARCA_LISTA.match(linha)
        and not MARCA_NUMERO.match(linha)
    )


# ---------- texto ----------

def _ler_texto(caminho: Path) -> Documento:
    texto = caminho.read_text(encoding="utf-8", errors="replace")
    grupos = re.split(r"\n\s*\n", texto.strip())
    blocos: list[Bloco] = []
    for grupo in grupos:
        linhas = grupo.splitlines()
        if len(linhas) == 1 and linhas[0].lstrip().startswith("#"):
            marcas = len(linhas[0]) - len(linhas[0].lstrip("#"))
            blocos.append(Bloco("titulo", linhas[0].lstrip("# ").strip(), nivel=min(marcas, 3)))
        elif len(linhas) == 1 and _parece_titulo(linhas[0]) and blocos is not None:
            blocos.append(Bloco("titulo", linhas[0].strip(), nivel=2))
        else:
            b = _bloco_de_linhas(linhas)
            if b:
                blocos.append(b)
    titulo, blocos = _tirar_titulo(blocos, caminho.stem.replace("-", " ").replace("_", " "))
    return Documento(titulo, blocos)


# ---------- PDF ----------

def _ler_pdf(caminho: Path, pasta_trabalho: Path) -> Documento:
    try:
        import pymupdf as fitz  # PyMuPDF
    except ImportError as erro:
        raise RuntimeError("Para ler PDF instale: pip install pymupdf") from erro

    pdf = fitz.open(caminho)
    paginas = [p.get_text("dict")["blocks"] for p in pdf]

    # tamanho de letra mais comum = corpo do texto; maiores = títulos
    pesos: Counter = Counter()
    for blocos_pg in paginas:
        for b in blocos_pg:
            for linha in b.get("lines", []):
                for s in linha["spans"]:
                    pesos[round(s["size"])] += len(s["text"].strip())
    corpo = pesos.most_common(1)[0][0] if pesos else 12

    pasta_img = pasta_trabalho / "imagens"
    n_img = 0
    blocos: list[Bloco] = []
    for blocos_pg in paginas:
        for b in blocos_pg:
            if b["type"] == 1:  # imagem
                if min(b["bbox"][2] - b["bbox"][0], b["bbox"][3] - b["bbox"][1]) < 40:
                    continue  # ícones e enfeites minúsculos
                n_img += 1
                pasta_img.mkdir(parents=True, exist_ok=True)
                nome = f"img-{n_img:03d}.{b.get('ext', 'png')}"
                (pasta_img / nome).write_bytes(b["image"])
                blocos.append(Bloco("imagem", arquivo=f"imagens/{nome}"))
                continue
            # o PDF às vezes junta num bloco linhas de tamanhos diferentes (ex.: cabeçalho pequeno + título grande);
            # separa em grupos de linhas com o mesmo tamanho e o mesmo peso
            grupos: list[list[tuple[str, float, bool]]] = []
            for linha in b.get("lines", []):
                t = "".join(s["text"] for s in linha["spans"]).strip()
                if not t:
                    continue
                tam = max(s["size"] for s in linha["spans"])
                negrito = all(s["flags"] & 16 or "bold" in s["font"].lower() for s in linha["spans"] if s["text"].strip())
                if grupos and abs(grupos[-1][-1][1] - tam) <= 0.15 * tam and grupos[-1][-1][2] == negrito:
                    grupos[-1].append((t, tam, negrito))
                else:
                    grupos.append([(t, tam, negrito)])
            for grupo in grupos:
                linhas = [g[0] for g in grupo]
                if re.fullmatch(r"\d{1,3}", " ".join(linhas)):
                    continue  # número de página
                texto = " ".join(linhas)
                tam = max(g[1] for g in grupo)
                if len(texto) <= 120 and tam >= corpo * 1.2:
                    blocos.append(Bloco("titulo", texto, nivel=1 if tam >= corpo * 1.6 else 2))
                elif len(texto) <= 80 and all(g[2] for g in grupo) and not texto.endswith("."):
                    blocos.append(Bloco("titulo", texto, nivel=3))
                else:
                    bloco = _bloco_de_linhas(linhas)
                    if bloco:
                        if bloco.tipo == "paragrafo":
                            bloco.texto = re.sub(r"(\w)- (?=[a-zà-ú])", r"\1", bloco.texto)
                        blocos.append(bloco)

    titulo, blocos = _tirar_titulo(blocos, caminho.stem)
    doc = Documento(titulo, blocos)
    if not any(b.tipo == "paragrafo" or b.tipo == "lista" for b in blocos):
        doc.pendencias.append(
            "Este PDF parece ser uma imagem escaneada (sem texto). É preciso aplicar OCR antes."
        )
    return doc


# ---------- PowerPoint ----------

def _ler_pptx(caminho: Path, pasta_trabalho: Path) -> Documento:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as erro:
        raise RuntimeError("Para ler PowerPoint instale: pip install python-pptx") from erro

    apresentacao = Presentation(str(caminho))
    pasta_img = pasta_trabalho / "imagens"
    n_img = 0
    blocos: list[Bloco] = []
    for slide in apresentacao.slides:
        titulo_slide = slide.shapes.title
        if titulo_slide is not None and titulo_slide.text_frame.text.strip():
            blocos.append(Bloco("titulo", titulo_slide.text_frame.text.strip(), nivel=2))
        formas = sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0))
        for forma in formas:
            if titulo_slide is not None and forma.shape_id == titulo_slide.shape_id:
                continue
            if forma.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n_img += 1
                pasta_img.mkdir(parents=True, exist_ok=True)
                nome = f"img-{n_img:03d}.{forma.image.ext}"
                (pasta_img / nome).write_bytes(forma.image.blob)
                alt = forma._element.xpath(".//p:cNvPr")[0].get("descr") or None
                if alt and NOME_DE_ARQUIVO.match(alt.strip()):
                    alt = None  # o PowerPoint às vezes grava o nome do arquivo como "descrição"
                blocos.append(Bloco("imagem", arquivo=f"imagens/{nome}", alt=alt))
            elif getattr(forma, "has_table", False) and forma.has_table:
                linhas = [" — ".join(c.text.strip() for c in l.cells) for l in forma.table.rows]
                blocos.append(Bloco("lista", itens=[l for l in linhas if l.strip(" —")]))
            elif forma.has_text_frame:
                paragrafos = [p.text.strip() for p in forma.text_frame.paragraphs if p.text.strip()]
                if len(paragrafos) == 1:
                    blocos.append(Bloco("paragrafo", texto=paragrafos[0]))
                elif paragrafos:
                    b = _bloco_de_linhas(paragrafos)
                    # vários parágrafos no mesmo quadro de slide são tópicos, mesmo sem marcador
                    blocos.append(b if b and b.tipo != "paragrafo" else Bloco("lista", itens=paragrafos))
    primeiro = apresentacao.core_properties.title or caminho.stem
    return Documento(primeiro, blocos)


# ---------- vídeo e áudio ----------

def _ler_midia(caminho: Path, modelo: str) -> Documento:
    """Transcreve o áudio localmente com faster-whisper (sem internet depois de baixar o modelo)."""
    try:
        from faster_whisper import WhisperModel
    except ImportError as erro:
        raise RuntimeError(
            "Para transcrever vídeo/áudio instale: pip install faster-whisper"
        ) from erro

    whisper = WhisperModel(modelo, device="cpu", compute_type="int8")
    segmentos, _ = whisper.transcribe(str(caminho), language="pt", vad_filter=True)

    # agrupa frases em parágrafos: nova pausa longa (>1,5 s) ou ~5 frases
    paragrafos, atual, fim_anterior = [], [], 0.0
    for seg in segmentos:
        if atual and (seg.start - fim_anterior > 1.5 or len(atual) >= 5):
            paragrafos.append(" ".join(atual))
            atual = []
        atual.append(seg.text.strip())
        fim_anterior = seg.end
    if atual:
        paragrafos.append(" ".join(atual))

    doc = Documento(caminho.stem, [Bloco("paragrafo", texto=p) for p in paragrafos])
    doc.pendencias.append(
        "O texto veio da transcrição automática do áudio: o professor precisa conferir nomes, "
        "números e termos técnicos. Imagens e slides mostrados no vídeo não foram lidos."
    )
    return doc
