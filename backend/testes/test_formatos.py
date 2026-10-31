"""Testes dos leitores de PDF e PowerPoint. Montam os arquivos na hora, sem precisar de exemplos.

Só rodam se pymupdf / python-pptx estiverem instalados (pip install -r requirements.txt).
"""
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from adaptador import cache, leitores, pipeline, transformacoes as t

try:
    import pymupdf
except ImportError:
    pymupdf = None
try:
    import pptx
except ImportError:
    pptx = None


def _png() -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 120, 90), False)
    pix.clear_with(180)
    return pix.tobytes("png")


@unittest.skipIf(pymupdf is None, "pymupdf não instalado")
class TestPDF(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        pdf = pymupdf.open()
        pg = pdf.new_page()
        pg.insert_text((72, 80), "O ciclo da água", fontsize=26)
        pg.insert_text((72, 130), "As etapas", fontsize=18)
        pg.insert_textbox(pymupdf.Rect(72, 150, 520, 260), (
            "A água da Terra muda de lugar e de forma o tempo todo. Ela evapora, forma nuvens, "
            "cai como chuva e volta aos rios e ao mar. Esse caminho se repete sempre."), fontsize=11)
        pg.insert_textbox(pymupdf.Rect(72, 270, 520, 340), (
            "Pinte a nuvem de azul e circule em vermelho o sol."), fontsize=11)
        pg.insert_image(pymupdf.Rect(72, 360, 192, 450), stream=_png())
        pg.insert_text((300, 800), "1", fontsize=10)  # número de página
        self.caminho = self.dir / "aula.pdf"
        pdf.save(self.caminho)

    def tearDown(self):
        self.tmp.cleanup()

    def test_le_titulos_paragrafos_e_imagem(self):
        doc = leitores.ler(self.caminho, self.dir / "trab")
        self.assertEqual(doc.titulo, "O ciclo da água")
        tipos = [b.tipo for b in doc.blocos]
        self.assertEqual(tipos[0], "titulo")           # "As etapas"
        self.assertIn("paragrafo", tipos)
        self.assertIn("imagem", tipos)
        self.assertNotIn("1", [b.texto for b in doc.blocos])  # número de página ignorado
        img = next(b for b in doc.blocos if b.tipo == "imagem")
        self.assertTrue((self.dir / "trab" / img.arquivo).exists())

    def test_pipeline_gera_pendencia_de_imagem_e_simbolos(self):
        antigo, cache.PASTA = cache.PASTA, self.dir / "cache"
        try:
            r = pipeline.adaptar(self.caminho, {"necessidades": ["baixa_visao"]}, self.dir / "saida")
        finally:
            cache.PASTA = antigo
        self.assertFalse(r["pode_publicar"])           # imagem sem descrição (P4.1)
        html = (Path(r["pasta"]) / "ampliada.html").read_text(encoding="utf-8")
        self.assertIn("azul (■)", html)
        self.assertTrue((Path(r["pasta"]) / "imagens").exists())

    def test_cabecalho_pequeno_nao_gruda_no_titulo(self):
        pdf = pymupdf.open()
        pg = pdf.new_page()
        pg.insert_text((72, 70), "Ciências - 5º ano", fontsize=9)
        pg.insert_text((72, 100), "O ciclo da água", fontsize=22)     # o PDF junta as duas linhas num bloco
        pg.insert_textbox(pymupdf.Rect(72, 120, 520, 220), "A água muda de estado o tempo todo. " * 3, fontsize=11)
        f = self.dir / "cab.pdf"
        pdf.save(f)
        doc = leitores.ler(f, self.dir / "t3")
        self.assertEqual(doc.titulo, "O ciclo da água")
        self.assertEqual((doc.blocos[0].tipo, doc.blocos[0].texto), ("paragrafo", "Ciências - 5º ano"))

    def test_marcador_do_word_vira_lista(self):
        marca = chr(0xF0B7)
        b = leitores._bloco_de_linhas([marca + " Evaporação", marca + " Condensação"])
        self.assertEqual((b.tipo, b.itens), ("lista", ["Evaporação", "Condensação"]))

    def test_pdf_escaneado_avisa(self):
        pdf = pymupdf.open()
        pg = pdf.new_page()
        pg.insert_image(pymupdf.Rect(50, 50, 500, 700), stream=_png())
        f = self.dir / "scan.pdf"
        pdf.save(f)
        doc = leitores.ler(f, self.dir / "trab2")
        self.assertTrue(any("OCR" in p for p in doc.pendencias))


@unittest.skipIf(pptx is None or pymupdf is None, "python-pptx/pymupdf não instalados")
class TestPPTX(unittest.TestCase):
    def setUp(self):
        from pptx import Presentation
        from pptx.util import Inches
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        p = Presentation()
        s1 = p.slides.add_slide(p.slide_layouts[1])      # título + conteúdo
        s1.shapes.title.text = "Evaporação"
        s1.placeholders[1].text_frame.text = "O sol aquece a água"
        s1.placeholders[1].text_frame.add_paragraph().text = "Ela sobe como vapor"
        s2 = p.slides.add_slide(p.slide_layouts[5])      # só título
        s2.shapes.title.text = "Atividade"
        (self.dir / "i.png").write_bytes(_png())
        s2.shapes.add_picture(str(self.dir / "i.png"), Inches(1), Inches(2))
        tab = s2.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
        tab.cell(0, 0).text, tab.cell(0, 1).text = "Estado", "Exemplo"
        tab.cell(1, 0).text, tab.cell(1, 1).text = "Gasoso", "Vapor"
        self.caminho = self.dir / "aula.pptx"
        p.save(self.caminho)

    def tearDown(self):
        self.tmp.cleanup()

    def test_le_slides(self):
        doc = leitores.ler(self.caminho, self.dir / "trab")
        textos = [(b.tipo, b.texto or b.itens) for b in doc.blocos]
        self.assertIn(("titulo", "Evaporação"), textos)
        self.assertIn(("lista", ["O sol aquece a água", "Ela sobe como vapor"]), textos)
        self.assertIn(("lista", ["Estado — Exemplo", "Gasoso — Vapor"]), textos)
        img = next(b for b in doc.blocos if b.tipo == "imagem")
        self.assertIsNone(img.alt)
        self.assertTrue((self.dir / "trab" / img.arquivo).exists())

    def test_imagem_sem_alt_vira_pendencia(self):
        doc = leitores.ler(self.caminho, self.dir / "trab")
        self.assertTrue(t.imagens_com_descricao(doc).pendencias)


if __name__ == "__main__":
    unittest.main()
